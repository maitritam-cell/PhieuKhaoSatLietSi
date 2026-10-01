import io, json, os
from http.server import BaseHTTPRequestHandler
from docx import Document

TEMPLATE = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'template', 'Mẫu 02.docx')

def replace_first(paragraph, label, value):
    value=(value or '').strip()
    if not value: return False
    text=paragraph.text
    if label not in text: return False
    new=text.replace(label, label+' '+value, 1)
    for r in list(paragraph.runs): r._element.getparent().remove(r._element)
    run=paragraph.add_run(new)
    run.font.name='Times New Roman'
    return True

def fill_document(data):
    doc=Document(TEMPLATE)
    ps=doc.paragraphs
    mapping={
      4:[('Họ và tên:',data.get('rep_name'))],
      5:[('Ngày tháng năm sinh:',data.get('rep_dob')),('Giới tính:',data.get('rep_gender'))],
      6:[('Số ĐDCN',data.get('rep_id')),('Ngày cấp',data.get('rep_issue_date')),('Nơi cấp:',data.get('rep_issue_place'))],
      7:[('Quê quán:',data.get('rep_hometown'))],
      8:[('Nơi thường trú:',data.get('rep_address'))],
      9:[('Số điện thoại:',data.get('rep_phone'))],
      12:[('Mã hồ sơ Bộ quản lý:',data.get('ministry_file')),('Mã hồ sơ tỉnh quản lý:',data.get('province_file'))],
      13:[('Họ và tên liệt sĩ:',data.get('martyr_name')),('Bí danh:',data.get('martyr_alias'))],
      14:[('Ngày tháng năm sinh:',data.get('martyr_dob')),('Giới tính:',data.get('martyr_gender'))],
      15:[('Quê quán:',data.get('martyr_hometown'))],
      16:[('Cấp bậc, chức vụ khi hy sinh:',data.get('martyr_rank'))],
      17:[('Cơ quan, đơn vị khi hy sinh:',data.get('martyr_unit'))],
      18:[('Ngày tháng năm hy sinh:',data.get('martyr_death_date'))],
      19:[('Nơi hy sinh (nếu có):',data.get('martyr_death_place'))],
      20:[('Nơi an táng ban đầu:',data.get('burial_place'))],
      21:[('Bằng Tổ quốc ghi công số',data.get('certificate_no')),('Quyết định số',data.get('decision_no')),('ngày',data.get('decision_date'))],
      22:[('Con ông:',data.get('father'))],
      23:[('Con bà:',data.get('mother'))],
      24:[('Vợ:',data.get('wife'))]
    }
    for idx, items in mapping.items():
        if idx < len(ps):
            p=ps[idx]
            for label,value in items:
                if value: replace_first(p,label,value)
    if len(doc.tables) >= 3:
        table=doc.tables[2]
        relatives=data.get('relatives') or []
        for i, rel in enumerate(relatives[:10]):
            if i+2 >= len(table.rows): break
            row=table.rows[i+2]
            vals=[rel.get('id',''),rel.get('name',''),rel.get('dob',''),rel.get('gender',''),
                  rel.get('father',''),rel.get('mother',''),rel.get('address',''),rel.get('status',''),rel.get('signature','')]
            for col,val in zip([2,3,4,5,6,7,8,9,10], vals):
                row.cells[col].text=str(val or '')
    out=io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out.getvalue()

class handler(BaseHTTPRequestHandler):
    def do_POST(self):
        if self.path.rstrip('/') != '/api/generate':
            self.send_error(404)
            return
        try:
            n=int(self.headers.get('Content-Length','0'))
            data=json.loads(self.rfile.read(n).decode('utf-8'))
            content=fill_document(data)
            filename='Phieu_khao_sat_liet_si.docx'
            self.send_response(200)
            self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            self.send_header('Content-Disposition',f'attachment; filename="{filename}"')
            self.send_header('Content-Length',str(len(content)))
            self.end_headers()
            self.wfile.write(content)
        except Exception as e:
            body=json.dumps({'error':str(e)},ensure_ascii=False).encode('utf-8')
            self.send_response(500)
            self.send_header('Content-Type','application/json; charset=utf-8')
            self.send_header('Content-Length',str(len(body)))
            self.end_headers()
            self.wfile.write(body)
